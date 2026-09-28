# -*- coding: utf-8 -*-

"""
YOUTUBE COMMENT SCANNER - GUI
==============================

Chức năng:
1. Dán hàng loạt URL video YouTube vào hộp thoại
2. Nhập nhiều keyword vào hộp thoại
3. Quét:
   - Comment chính
   - Toàn bộ reply có thể truy cập
4. Xuất Excel:
   - ALL_COMMENTS
   - MATCHES
   - AUTHORS
   - VIDEO_SUMMARY
   - VIDEOS
   - SCAN_INFO

API:
    api.py
    └── API_KEY = "YOUR_API_KEY"

Cài:
    pip install requests pandas openpyxl tqdm
"""

import os
import re
import time
import unicodedata
import threading
import queue
import tkinter as tk

from tkinter import ttk, messagebox, filedialog

from datetime import datetime

import requests
import pandas as pd

from api import API_KEY

# ============================================================
# CONFIG
# ============================================================

API_BASE = "https://www.googleapis.com/youtube/v3"

OUTPUT_FILE = "YOUTUBE_COMMENT_SCAN.xlsx"

REQUEST_DELAY = 0.05

MAX_RETRY = 5


# ============================================================
# GLOBAL STOP
# ============================================================

STOP_SCAN = False


# ============================================================
# GUI LOG QUEUE
# ============================================================

log_queue = queue.Queue()


def gui_log(text=""):
    """
    Gửi log từ thread scan sang GUI.
    """
    log_queue.put(str(text))


def process_logs():
    """
    Cập nhật log lên GUI.
    """

    try:

        while True:

            msg = log_queue.get_nowait()

            log_text.insert(
                tk.END,
                msg + "\n"
            )

            log_text.see(
                tk.END
            )

    except queue.Empty:
        pass

    root.after(
        100,
        process_logs
    )


# ============================================================
# BASIC HELPERS
# ============================================================

def normalize_text(text):

    if not text:
        return ""

    text = str(text).strip().lower()

    text = unicodedata.normalize(
        "NFD",
        text
    )

    text = "".join(
        c
        for c in text
        if unicodedata.category(c) != "Mn"
    )

    text = text.replace(
        "đ",
        "d"
    )

    return text


def clean_text(text):

    if text is None:
        return ""

    text = str(text)

    text = re.sub(
        r"\s+",
        " ",
        text
    )

    return text.strip()


# ============================================================
# VIDEO ID
# ============================================================

def extract_video_id(url):

    if not url:
        return None

    url = url.strip()

    # watch?v=
    m = re.search(
        r"[?&]v=([A-Za-z0-9_-]{11})",
        url
    )

    if m:
        return m.group(1)

    # youtu.be/
    m = re.search(
        r"youtu\.be/([A-Za-z0-9_-]{11})",
        url
    )

    if m:
        return m.group(1)

    # shorts/
    m = re.search(
        r"/shorts/([A-Za-z0-9_-]{11})",
        url
    )

    if m:
        return m.group(1)

    # live/
    m = re.search(
        r"/live/([A-Za-z0-9_-]{11})",
        url
    )

    if m:
        return m.group(1)

    # fallback
    m = re.search(
        r"\b([A-Za-z0-9_-]{11})\b",
        url
    )

    if m:
        return m.group(1)

    return None


# ============================================================
# API REQUEST
# ============================================================

def api_get(endpoint, params):

    url = API_BASE + "/" + endpoint

    for attempt in range(
        1,
        MAX_RETRY + 1
    ):

        if STOP_SCAN:
            return None

        try:

            r = requests.get(
                url,
                params=params,
                timeout=30
            )

            if r.status_code == 200:
                return r.json()

            try:

                error_json = r.json()

                error_message = (
                    error_json
                    .get("error", {})
                    .get("message", "")
                )

            except Exception:

                error_message = r.text[:500]

            # ------------------------------------------------
            # 403
            # ------------------------------------------------

            if r.status_code == 403:

                gui_log("")
                gui_log(
                    "[X] YouTube API 403"
                )

                gui_log(
                    "[X] " + error_message
                )

                if (
                    "quota"
                    in error_message.lower()
                ):

                    gui_log(
                        "[!] API quota có thể đã hết."
                    )

                    return None

            # ------------------------------------------------
            # RETRY
            # ------------------------------------------------

            if r.status_code in [
                429,
                500,
                502,
                503,
                504
            ]:

                wait = attempt * 2

                gui_log(
                    f"[!] HTTP {r.status_code}"
                    f" -> retry {wait}s..."
                )

                time.sleep(
                    wait
                )

                continue

            gui_log(
                f"[X] API error "
                f"{r.status_code}: "
                f"{error_message}"
            )

            return None

        except requests.RequestException as e:

            wait = attempt * 2

            gui_log(
                f"[!] Request error: {e}"
            )

            if attempt < MAX_RETRY:

                time.sleep(
                    wait
                )

    return None


# ============================================================
# GET VIDEO INFO
# ============================================================

def get_video_info(video_id):

    data = api_get(
        "videos",
        {
            "part":
                "snippet,statistics",

            "id":
                video_id,

            "key":
                API_KEY
        }
    )

    if not data:
        return None

    items = data.get(
        "items",
        []
    )

    if not items:
        return None

    item = items[0]

    snippet = item.get(
        "snippet",
        {}
    )

    stats = item.get(
        "statistics",
        {}
    )

    return {

        "video_id":
            video_id,

        "video_url":
            f"https://www.youtube.com/watch?v={video_id}",

        "video_title":
            snippet.get(
                "title",
                ""
            ),

        "channel_title":
            snippet.get(
                "channelTitle",
                ""
            ),

        "channel_id":
            snippet.get(
                "channelId",
                ""
            ),

        "published_at":
            snippet.get(
                "publishedAt",
                ""
            ),

        "view_count":
            stats.get(
                "viewCount",
                ""
            ),

        "like_count":
            stats.get(
                "likeCount",
                ""
            ),

        "comment_count":
            stats.get(
                "commentCount",
                ""
            )
    }


# ============================================================
# COMMENT OBJECT
# ============================================================

def parse_comment(
        item,
        video_info,
        comment_type,
        parent_id=""
):

    snippet = item.get(
        "snippet",
        {}
    )

    author_channel_id = (
        snippet
        .get(
            "authorChannelId",
            {}
        )
        .get(
            "value",
            ""
        )
    )

    author_name = snippet.get(
        "authorDisplayName",
        ""
    )

    author_channel_url = ""

    if author_channel_id:

        author_channel_url = (
            "https://www.youtube.com/channel/"
            + author_channel_id
        )

    comment_id = item.get(
        "id",
        ""
    )

    text_original = snippet.get(
        "textOriginal",
        ""
    )

    text_display = snippet.get(
        "textDisplay",
        ""
    )

    published_at = snippet.get(
        "publishedAt",
        ""
    )

    updated_at = snippet.get(
        "updatedAt",
        ""
    )

    like_count = snippet.get(
        "likeCount",
        0
    )

    return {

        "video_id":
            video_info.get(
                "video_id",
                ""
            ),

        "video_title":
            video_info.get(
                "video_title",
                ""
            ),

        "video_url":
            video_info.get(
                "video_url",
                ""
            ),

        "video_channel":
            video_info.get(
                "channel_title",
                ""
            ),

        "video_channel_id":
            video_info.get(
                "channel_id",
                ""
            ),

        "video_published_at":
            video_info.get(
                "published_at",
                ""
            ),

        "comment_type":
            comment_type,

        "comment_id":
            comment_id,

        "parent_id":
            parent_id,

        "published_at":
            published_at,

        "updated_at":
            updated_at,

        "like_count":
            like_count,

        "author_name":
            author_name,

        "author_channel_id":
            author_channel_id,

        "author_channel_url":
            author_channel_url,

        "comment_text":
            clean_text(
                text_original
            ),

        "comment_text_display":
            clean_text(
                text_display
            )
    }


# ============================================================
# GET ALL REPLIES
# ============================================================

def get_all_replies(
        parent_comment_id,
        video_info,
        rows
):

    next_page_token = None

    while True:

        if STOP_SCAN:
            break

        params = {

            "part":
                "snippet",

            "parentId":
                parent_comment_id,

            "maxResults":
                100,

            "key":
                API_KEY
        }

        if next_page_token:

            params[
                "pageToken"
            ] = next_page_token

        data = api_get(
            "comments",
            params
        )

        if not data:
            break

        items = data.get(
            "items",
            []
        )

        for item in items:

            if STOP_SCAN:
                break

            row = parse_comment(
                item=item,
                video_info=video_info,
                comment_type="REPLY",
                parent_id=parent_comment_id
            )

            rows.append(
                row
            )

        next_page_token = data.get(
            "nextPageToken"
        )

        if not next_page_token:
            break

        time.sleep(
            REQUEST_DELAY
        )


# ============================================================
# GET TOP COMMENTS
# ============================================================

def get_all_comments(
        video_info,
        rows,
        stats
):

    video_id = video_info[
        "video_id"
    ]

    next_page_token = None

    top_comment_count = 0
    reply_count = 0

    while True:

        if STOP_SCAN:
            break

        params = {

            "part":
                "snippet,replies",

            "videoId":
                video_id,

            "maxResults":
                100,

            "order":
                "time",

            "textFormat":
                "plainText",

            "key":
                API_KEY
        }

        if next_page_token:

            params[
                "pageToken"
            ] = next_page_token

        data = api_get(
            "commentThreads",
            params
        )

        if not data:
            break

        items = data.get(
            "items",
            []
        )

        for item in items:

            if STOP_SCAN:
                break

            top_comment_count += 1

            top = (
                item
                .get(
                    "snippet",
                    {}
                )
                .get(
                    "topLevelComment",
                    {}
                )
            )

            if not top:
                continue

            row = parse_comment(
                item=top,
                video_info=video_info,
                comment_type="TOP_LEVEL",
                parent_id=""
            )

            rows.append(
                row
            )

            total_reply_count = (
                item
                .get(
                    "snippet",
                    {}
                )
                .get(
                    "totalReplyCount",
                    0
                )
            )

            if total_reply_count > 0:

                before = len(
                    rows
                )

                get_all_replies(
                    parent_comment_id=
                    top.get(
                        "id",
                        ""
                    ),

                    video_info=
                    video_info,

                    rows=
                    rows
                )

                after = len(
                    rows
                )

                reply_count += (
                    after - before
                )

        next_page_token = data.get(
            "nextPageToken"
        )

        if not next_page_token:
            break

        time.sleep(
            REQUEST_DELAY
        )

    stats[
        "top_level_comments"
    ] = top_comment_count

    stats[
        "replies"
    ] = reply_count

    stats[
        "total_comments"
    ] = (
        top_comment_count
        + reply_count
    )


# ============================================================
# KEYWORD MATCH
# ============================================================

def build_keyword_matches(
        rows,
        keywords
):

    if not keywords:
        return []

    normalized_keywords = []

    for kw in keywords:

        kw_clean = clean_text(
            kw
        )

        if not kw_clean:
            continue

        normalized_keywords.append(
            (
                kw_clean,
                normalize_text(
                    kw_clean
                )
            )
        )

    matches = []

    for row in rows:

        fields = {

            "AUTHOR_NAME":
                normalize_text(
                    row.get(
                        "author_name",
                        ""
                    )
                ),

            "AUTHOR_CHANNEL_ID":
                normalize_text(
                    row.get(
                        "author_channel_id",
                        ""
                    )
                ),

            "AUTHOR_CHANNEL_URL":
                normalize_text(
                    row.get(
                        "author_channel_url",
                        ""
                    )
                ),

            "COMMENT_TEXT":
                normalize_text(
                    row.get(
                        "comment_text",
                        ""
                    )
                )
        }

        matched_keywords = []

        matched_fields = []

        for (
            original_kw,
            norm_kw
        ) in normalized_keywords:

            if not norm_kw:
                continue

            for (
                field_name,
                field_value
            ) in fields.items():

                if norm_kw in field_value:

                    if (
                        original_kw
                        not in matched_keywords
                    ):

                        matched_keywords.append(
                            original_kw
                        )

                    if (
                        field_name
                        not in matched_fields
                    ):

                        matched_fields.append(
                            field_name
                        )

        if matched_keywords:

            new_row = dict(
                row
            )

            new_row[
                "matched_keywords"
            ] = " | ".join(
                matched_keywords
            )

            new_row[
                "matched_fields"
            ] = " | ".join(
                matched_fields
            )

            matches.append(
                new_row
            )

    return matches


# ============================================================
# AUTHORS
# ============================================================

def build_authors(rows):

    data = {}

    for row in rows:

        channel_id = row.get(
            "author_channel_id",
            ""
        )

        author_name = row.get(
            "author_name",
            ""
        )

        author_url = row.get(
            "author_channel_url",
            ""
        )

        key = (
            channel_id
            if channel_id
            else
            (
                author_name.lower(),
                author_url.lower()
            )
        )

        if key not in data:

            data[key] = {

                "author_name":
                    author_name,

                "author_channel_id":
                    channel_id,

                "author_channel_url":
                    author_url,

                "comment_count":
                    0,

                "videos_commented":
                    set(),

                "first_comment":
                    row.get(
                        "published_at",
                        ""
                    ),

                "last_comment":
                    row.get(
                        "published_at",
                        ""
                    )
            }

        obj = data[key]

        obj[
            "comment_count"
        ] += 1

        if row.get(
            "video_id"
        ):

            obj[
                "videos_commented"
            ].add(
                row["video_id"]
            )

        published = row.get(
            "published_at",
            ""
        )

        if published:

            if (
                not obj[
                    "first_comment"
                ]
                or
                published <
                obj[
                    "first_comment"
                ]
            ):

                obj[
                    "first_comment"
                ] = published

            if (
                not obj[
                    "last_comment"
                ]
                or
                published >
                obj[
                    "last_comment"
                ]
            ):

                obj[
                    "last_comment"
                ] = published

    result = []

    for obj in data.values():

        result.append({

            "author_name":
                obj[
                    "author_name"
                ],

            "author_channel_id":
                obj[
                    "author_channel_id"
                ],

            "author_channel_url":
                obj[
                    "author_channel_url"
                ],

            "comment_count":
                obj[
                    "comment_count"
                ],

            "videos_commented":
                len(
                    obj[
                        "videos_commented"
                    ]
                ),

            "first_comment":
                obj[
                    "first_comment"
                ],

            "last_comment":
                obj[
                    "last_comment"
                ]
        })

    return result


# ============================================================
# EXPORT EXCEL
# ============================================================

def export_excel(
        all_rows,
        matches,
        authors,
        video_summaries,
        videos,
        keywords,
        output_file
):

    gui_log("")
    gui_log(
        "[>] Đang xuất Excel..."
    )

    df_all = pd.DataFrame(
        all_rows
    )

    df_matches = pd.DataFrame(
        matches
    )

    df_authors = pd.DataFrame(
        authors
    )

    df_video_summary = pd.DataFrame(
        video_summaries
    )

    df_videos = pd.DataFrame(
        videos
    )

    scan_info = pd.DataFrame([
        {

            "scan_time":
                datetime.now().strftime(
                    "%Y-%m-%d %H:%M:%S"
                ),

            "keyword_count":
                len(
                    keywords
                ),

            "keywords":
                " | ".join(
                    keywords
                ),

            "videos":
                len(
                    videos
                ),

            "all_comments":
                len(
                    all_rows
                ),

            "matched_comments":
                len(
                    matches
                )
        }
    ])

    with pd.ExcelWriter(
        output_file,
        engine="openpyxl"
    ) as writer:

        df_all.to_excel(
            writer,
            sheet_name="ALL_COMMENTS",
            index=False
        )

        df_matches.to_excel(
            writer,
            sheet_name="MATCHES",
            index=False
        )

        df_authors.to_excel(
            writer,
            sheet_name="AUTHORS",
            index=False
        )

        df_video_summary.to_excel(
            writer,
            sheet_name="VIDEO_SUMMARY",
            index=False
        )

        df_videos.to_excel(
            writer,
            sheet_name="VIDEOS",
            index=False
        )

        scan_info.to_excel(
            writer,
            sheet_name="SCAN_INFO",
            index=False
        )

    # --------------------------------------------------------
    # FORMAT
    # --------------------------------------------------------

    try:

        from openpyxl import load_workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter

        wb = load_workbook(
            output_file
        )

        for ws in wb.worksheets:

            ws.freeze_panes = "A2"

            ws.auto_filter.ref = (
                ws.dimensions
            )

            for cell in ws[1]:

                cell.font = Font(
                    bold=True
                )

            for col_cells in ws.columns:

                max_length = 0

                for cell in col_cells:

                    try:

                        value = str(
                            cell.value
                            if cell.value is not None
                            else ""
                        )

                        max_length = max(
                            max_length,
                            len(value)
                        )

                    except Exception:
                        pass

                width = min(
                    max(
                        max_length + 2,
                        10
                    ),
                    60
                )

                ws.column_dimensions[
                    get_column_letter(
                        col_cells[0].column
                    )
                ].width = width

        wb.save(
            output_file
        )

    except Exception as e:

        gui_log(
            "[!] Không format được Excel: "
            + str(e)
        )

    gui_log(
        f"[OK] Đã xuất: {output_file}"
    )


# ============================================================
# SCAN FUNCTION
# ============================================================

def run_scan(
        video_urls,
        keywords
):

    global STOP_SCAN

    STOP_SCAN = False

    if not API_KEY:

        gui_log(
            "[X] API_KEY đang trống."
        )

        return

    # --------------------------------------------------------
    # VIDEO IDS
    # --------------------------------------------------------

    videos = []

    seen_ids = set()

    for url in video_urls:

        video_id = extract_video_id(
            url
        )

        if not video_id:

            gui_log(
                "[!] Không lấy được Video ID:"
            )

            gui_log(
                "    " + url
            )

            continue

        if video_id in seen_ids:

            gui_log(
                f"[!] Bỏ video trùng: "
                f"{video_id}"
            )

            continue

        seen_ids.add(
            video_id
        )

        videos.append({

            "video_id":
                video_id,

            "video_url":
                f"https://www.youtube.com/watch?v={video_id}",

            "input_url":
                url
        })

    if not videos:

        gui_log(
            "[X] Không có Video ID hợp lệ."
        )

        return

    # --------------------------------------------------------
    # CONFIG
    # --------------------------------------------------------

    gui_log("")
    gui_log(
        "=" * 70
    )

    gui_log(
        "CẤU HÌNH SCAN"
    )

    gui_log(
        "=" * 70
    )

    gui_log(
        f"Videos: {len(videos)}"
    )

    if keywords:

        gui_log(
            f"Keywords: {len(keywords)}"
        )

        for kw in keywords:

            gui_log(
                "  - " + kw
            )

    else:

        gui_log(
            "Keywords: KHÔNG CÓ"
        )

        gui_log(
            "=> Sẽ xuất TOÀN BỘ comment."
        )

    gui_log(
        "=" * 70
    )

    # --------------------------------------------------------
    # SCAN
    # --------------------------------------------------------

    all_rows = []

    video_summaries = []

    successful_videos = []

    total_videos = len(
        videos
    )

    for index, video in enumerate(
        videos,
        start=1
    ):

        if STOP_SCAN:
            break

        video_id = video[
            "video_id"
        ]

        gui_log("")
        gui_log(
            "-" * 70
        )

        gui_log(
            f"[{index}/{total_videos}]"
        )

        gui_log(
            f"Video ID: {video_id}"
        )

        # ----------------------------------------------------
        # VIDEO INFO
        # ----------------------------------------------------

        video_info = get_video_info(
            video_id
        )

        if not video_info:

            gui_log(
                "[X] Không lấy được "
                "thông tin video."
            )

            video_summaries.append({

                "video_id":
                    video_id,

                "video_url":
                    video["video_url"],

                "status":
                    "ERROR_VIDEO_INFO",

                "top_level_comments":
                    0,

                "replies":
                    0,

                "total_comments":
                    0
            })

            continue

        gui_log(
            "Title: "
            + video_info[
                "video_title"
            ]
        )

        gui_log(
            "Channel: "
            + video_info[
                "channel_title"
            ]
        )

        rows_before = len(
            all_rows
        )

        stats = {}

        try:

            get_all_comments(
                video_info=
                video_info,

                rows=
                all_rows,

                stats=
                stats
            )

        except Exception as e:

            gui_log(
                "[X] Lỗi scan: "
                + repr(e)
            )

            stats.setdefault(
                "top_level_comments",
                0
            )

            stats.setdefault(
                "replies",
                0
            )

            stats.setdefault(
                "total_comments",
                0
            )

        rows_after = len(
            all_rows
        )

        gui_log(
            f"[OK] Top-level: "
            f"{stats.get('top_level_comments', 0):,}"
        )

        gui_log(
            f"[OK] Replies: "
            f"{stats.get('replies', 0):,}"
        )

        gui_log(
            f"[OK] Tổng: "
            f"{stats.get('total_comments', 0):,}"
        )

        video_summaries.append({

            "video_id":
                video_id,

            "video_url":
                video_info[
                    "video_url"
                ],

            "video_title":
                video_info[
                    "video_title"
                ],

            "channel_title":
                video_info[
                    "channel_title"
                ],

            "channel_id":
                video_info[
                    "channel_id"
                ],

            "video_published_at":
                video_info[
                    "published_at"
                ],

            "api_comment_count":
                video_info[
                    "comment_count"
                ],

            "top_level_comments":
                stats.get(
                    "top_level_comments",
                    0
                ),

            "replies":
                stats.get(
                    "replies",
                    0
                ),

            "total_comments":
                stats.get(
                    "total_comments",
                    0
                ),

            "rows_added":
                rows_after - rows_before,

            "status":
                "OK"
        })

        successful_videos.append(
            video_id
        )

    # --------------------------------------------------------
    # MATCH
    # --------------------------------------------------------

    gui_log("")
    gui_log(
        "[>] Đang tìm keyword..."
    )

    matches = build_keyword_matches(
        all_rows,
        keywords
    )

    # --------------------------------------------------------
    # AUTHORS
    # --------------------------------------------------------

    gui_log(
        "[>] Đang tổng hợp AUTHORS..."
    )

    authors = build_authors(
        all_rows
    )

    # --------------------------------------------------------
    # EXPORT
    # --------------------------------------------------------

    output_file = os.path.join(
        os.getcwd(),
        OUTPUT_FILE
    )

    export_excel(
        all_rows=
        all_rows,

        matches=
        matches,

        authors=
        authors,

        video_summaries=
        video_summaries,

        videos=
        videos,

        keywords=
        keywords,

        output_file=
        output_file
    )

    # --------------------------------------------------------
    # RESULT
    # --------------------------------------------------------

    gui_log("")
    gui_log(
        "=" * 70
    )

    if STOP_SCAN:

        gui_log(
            "[!] ĐÃ DỪNG SCAN"
        )

    else:

        gui_log(
            "HOÀN TẤT"
        )

    gui_log(
        "=" * 70
    )

    gui_log(
        f"Videos scanned: "
        f"{len(successful_videos):,}"
    )

    gui_log(
        f"All comments: "
        f"{len(all_rows):,}"
    )

    gui_log(
        f"Matched comments: "
        f"{len(matches):,}"
    )

    gui_log(
        f"Unique authors: "
        f"{len(authors):,}"
    )

    gui_log("")
    gui_log(
        "OUTPUT:"
    )

    gui_log(
        os.path.abspath(
            output_file
        )
    )

    gui_log("")
    gui_log(
        "Sheets:"
    )

    gui_log(
        "1. ALL_COMMENTS"
    )

    gui_log(
        "2. MATCHES"
    )

    gui_log(
        "3. AUTHORS"
    )

    gui_log(
        "4. VIDEO_SUMMARY"
    )

    gui_log(
        "5. VIDEOS"
    )

    gui_log(
        "6. SCAN_INFO"
    )

    gui_log(
        "=" * 70
    )

    root.after(
        0,
        lambda: messagebox.showinfo(
            "Hoàn tất",
            "Đã quét xong!\n\n"
            f"Videos: {len(successful_videos):,}\n"
            f"Comments: {len(all_rows):,}\n"
            f"Matches: {len(matches):,}\n\n"
            f"File:\n{output_file}"
        )
    )


# ============================================================
# START SCAN
# ============================================================

def start_scan():

    global STOP_SCAN

    if not API_KEY:

        messagebox.showerror(
            "API KEY",
            "API_KEY đang trống.\n\n"
            "Kiểm tra file api.py"
        )

        return

    # --------------------------------------------------------
    # GET VIDEO TEXT
    # --------------------------------------------------------

    raw_videos = (
        video_text.get(
            "1.0",
            tk.END
        )
        .strip()
    )

    if not raw_videos:

        messagebox.showwarning(
            "Thiếu dữ liệu",
            "Bạn chưa dán link video."
        )

        return

    video_urls = [

        x.strip()

        for x in raw_videos.splitlines()

        if x.strip()
    ]

    # --------------------------------------------------------
    # GET KEYWORDS
    # --------------------------------------------------------

    raw_keywords = (
        keyword_text.get(
            "1.0",
            tk.END
        )
        .strip()
    )

    keywords = [

        x.strip()

        for x in raw_keywords.splitlines()

        if x.strip()
    ]

    # --------------------------------------------------------
    # CLEAR LOG
    # --------------------------------------------------------

    log_text.delete(
        "1.0",
        tk.END
    )

    STOP_SCAN = False

    # --------------------------------------------------------
    # BUTTON
    # --------------------------------------------------------

    scan_button.config(
        state=tk.DISABLED
    )

    stop_button.config(
        state=tk.NORMAL
    )

    # --------------------------------------------------------
    # THREAD
    # --------------------------------------------------------

    thread = threading.Thread(
        target=run_scan,
        args=(
            video_urls,
            keywords
        ),
        daemon=True
    )

    thread.start()


# ============================================================
# STOP
# ============================================================

def stop_scan():

    global STOP_SCAN

    STOP_SCAN = True

    gui_log("")
    gui_log(
        "[!] Đang yêu cầu dừng..."
    )

    stop_button.config(
        state=tk.DISABLED
    )


# ============================================================
# CLEAR
# ============================================================

def clear_all():

    global STOP_SCAN

    STOP_SCAN = True

    video_text.delete(
        "1.0",
        tk.END
    )

    keyword_text.delete(
        "1.0",
        tk.END
    )

    log_text.delete(
        "1.0",
        tk.END
    )


# ============================================================
# OPEN OUTPUT
# ============================================================

def open_output_folder():

    path = os.path.abspath(
        OUTPUT_FILE
    )

    folder = os.path.dirname(
        path
    )

    try:

        if os.name == "nt":

            os.startfile(
                folder
            )

        elif sys.platform == "darwin":

            os.system(
                f'open "{folder}"'
            )

        else:

            os.system(
                f'xdg-open "{folder}"'
            )

    except Exception as e:

        messagebox.showerror(
            "Error",
            str(e)
        )


# ============================================================
# GUI
# ============================================================

root = tk.Tk()

root.title(
    "YouTube Comment Scanner"
)

root.geometry(
    "1050x850"
)

root.minsize(
    850,
    650
)


# ============================================================
# STYLE
# ============================================================

style = ttk.Style()

try:
    style.theme_use(
        "clam"
    )
except:
    pass


# ============================================================
# TITLE
# ============================================================

title_label = ttk.Label(
    root,
    text="YOUTUBE COMMENT SCANNER",
    font=(
        "Arial",
        18,
        "bold"
    )
)

title_label.pack(
    pady=12
)


subtitle = ttk.Label(
    root,
    text=(
        "Dán hàng loạt link video + nhập keyword "
        "→ quét comment và xuất Excel"
    ),
    font=(
        "Arial",
        10
    )
)

subtitle.pack(
    pady=(0, 10)
)


# ============================================================
# MAIN FRAME
# ============================================================

main_frame = ttk.Frame(
    root
)

main_frame.pack(
    fill=tk.BOTH,
    expand=True,
    padx=15,
    pady=5
)


# ============================================================
# LEFT / VIDEO
# ============================================================

video_frame = ttk.LabelFrame(
    main_frame,
    text=" 1. LINK VIDEO YOUTUBE ",
    padding=10
)

video_frame.pack(
    fill=tk.BOTH,
    expand=True,
    pady=5
)

video_hint = ttk.Label(
    video_frame,
    text=(
        "Dán nhiều link, mỗi dòng 1 video. "
        "Có thể dán trực tiếp Ctrl+V."
    )
)

video_hint.pack(
    anchor="w",
    pady=(0, 5)
)

video_text = tk.Text(
    video_frame,
    height=10,
    wrap=tk.NONE,
    font=(
        "Consolas",
        10
    )
)

video_text.pack(
    fill=tk.BOTH,
    expand=True
)


# ============================================================
# KEYWORDS
# ============================================================

keyword_frame = ttk.LabelFrame(
    main_frame,
    text=" 2. TỪ KHÓA ",
    padding=10
)

keyword_frame.pack(
    fill=tk.BOTH,
    expand=True,
    pady=5
)

keyword_hint = ttk.Label(
    keyword_frame,
    text=(
        "Mỗi dòng 1 keyword. Có thể để trống "
        "nếu muốn lấy toàn bộ comment."
    )
)

keyword_hint.pack(
    anchor="w",
    pady=(0, 5)
)

keyword_text = tk.Text(
    keyword_frame,
    height=7,
    wrap=tk.NONE,
    font=(
        "Consolas",
        10
    )
)

keyword_text.pack(
    fill=tk.BOTH,
    expand=True
)


# ============================================================
# BUTTONS
# ============================================================

button_frame = ttk.Frame(
    root
)

button_frame.pack(
    fill=tk.X,
    padx=15,
    pady=10
)


scan_button = ttk.Button(
    button_frame,
    text="▶  QUÉT COMMENT",
    command=start_scan
)

scan_button.pack(
    side=tk.LEFT,
    padx=5
)


stop_button = ttk.Button(
    button_frame,
    text="■  DỪNG",
    command=stop_scan,
    state=tk.DISABLED
)

stop_button.pack(
    side=tk.LEFT,
    padx=5
)


clear_button = ttk.Button(
    button_frame,
    text="XÓA",
    command=clear_all
)

clear_button.pack(
    side=tk.LEFT,
    padx=5
)


open_button = ttk.Button(
    button_frame,
    text="MỞ THƯ MỤC OUTPUT",
    command=open_output_folder
)

open_button.pack(
    side=tk.RIGHT,
    padx=5
)


# ============================================================
# LOG
# ============================================================

log_frame = ttk.LabelFrame(
    root,
    text=" 3. TIẾN TRÌNH ",
    padding=8
)

log_frame.pack(
    fill=tk.BOTH,
    expand=True,
    padx=15,
    pady=(0, 15)
)

log_text = tk.Text(
    log_frame,
    height=12,
    wrap=tk.WORD,
    bg="#111111",
    fg="#00ff66",
    insertbackground="white",
    font=(
        "Consolas",
        9
    )
)

log_text.pack(
    fill=tk.BOTH,
    expand=True
)


# ============================================================
# START LOG PROCESS
# ============================================================

root.after(
    100,
    process_logs
)


# ============================================================
# START GUI
# ============================================================

root.mainloop()